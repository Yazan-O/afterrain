"""Transfer rivers: Farleigh Hungerford (River Frome) and Alewife Brook (Massachusetts)."""
from __future__ import annotations

import csv
import datetime as dt
import io
import statistics as st
from typing import Any

import numpy as np
import openpyxl

from . import bath, config, model, sources
from .fetch import Fetcher

BROWSER = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/128 Safari/537.36"}
MASSDEP_HEADERS = {**BROWSER, "Accept": "application/json, text/plain, */*",
                   "Referer": "https://eeaonline.eea.state.ma.us/portal/dep/cso-data-portal/"}
ALEWIFE_WINDOW = (dt.datetime(2022, 8, 5), dt.datetime(2025, 10, 24))


def classified(samples: list[sources.Sample]) -> list[tuple[dt.datetime, float, bool | None]]:
    """(time, published value, label) per sample; the label is None when the qualifier leaves it undecided."""
    return [(x.t, x.ecoli, x.over_900) for x in samples]


def spill_backtest(samples: list[tuple[dt.datetime, float, bool | None]],
                   events: list[tuple[dt.datetime, dt.datetime]], hours: float = config.WINDOW_H) -> dict[str, Any]:
    """Samples warned when any spill overlaps the `hours` before them (positive overlap). A sample whose label the
    laboratory qualifier leaves undecided against 900 is dropped from every count and median."""
    rows = []
    for t, v, y in samples:
        if y is None:
            continue
        w0 = t - dt.timedelta(hours=hours)
        h = sum((min(e, t) - max(s, w0)).total_seconds() / 3600 for s, e in events if e > w0 and s < t)
        rows.append((v, y, h > 0))
    a = [(v, y) for v, y, w in rows if w]
    b = [(v, y) for v, y, w in rows if not w]
    return {"n": len(rows), "warned": len(a), "warned_exceed": sum(y for _, y in a),
            "warned_median": st.median(v for v, _ in a) if a else None, "not_warned": len(b),
            "not_warned_exceed": sum(y for _, y in b), "not_warned_median": st.median(v for v, _ in b) if b else None,
            "ambiguous_dropped": sum(y is None for _, _, y in samples)}


def farleigh(f: Fetcher, b: bath.Bath, portable: model.Fitted) -> dict[str, Any]:
    spills = sources.farleigh_spills(f)
    samples = sources.farleigh_samples(f)
    rule = spill_backtest(classified(samples), [(s.start, s.end) for s in spills])
    ovs = model.group_overflows(spills)
    bath.locate(b.net, ovs, config.FARLEIGH)
    site = model.Site(ovs, b.rain, sources.tellisford_flow(f))
    first = min(s.start for s in spills)
    used = [x for x in samples if x.t >= first and x.over_900 is not None]
    raws = [model.raw_features(site, x.t, portable.velocity) for x in used]
    y = np.array([bool(x.over_900) for x in used], int)
    p = portable.predict(raws)
    return {"rule": rule, "overflows": len(ovs), "overflow_names": sorted(o.name for o in ovs.values()),
            "samples_first": samples[0].local if samples else None,
            "samples_last": samples[-1].local if samples else None,
            "ambiguous": [{"time_local": x.local, "qualifier": x.operator, "ecoli_per_100ml": x.ecoli}
                          for x in samples if x.over_900 is None],
            "qualified": sum(x.operator != "" for x in samples),
            "y": y, "p": p,
            "missing_flow_days": sum(r["flow_ratio"] is None for r in raws)}


def _massdep_starts(path) -> set[tuple[str, dt.datetime]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb["Incidents"].iter_rows(values_only=True))
    head = rows[0]
    out = set()
    for r in rows[1:]:
        d = dict(zip(head, r))
        try:
            t = dt.datetime.strptime(d["IncidentDate"] + " " + (d["IncidentTime"] or "00:00"), "%m/%d/%Y %H:%M")
        except (TypeError, ValueError):
            continue
        out.add((d["OutfallId"], t))
    return out


def _mwra_starts(text: str) -> set[tuple[str, dt.datetime]]:
    out = set()
    for r in csv.DictReader(io.StringIO(text)):
        if r["CSO Number"] == "MWR003":
            try:
                out.add(("MWR003", dt.datetime.strptime(r["Start Time"], "%m/%d/%Y %I:%M %p")))
            except ValueError:
                pass
    return out


def _mystic_samples(path) -> list[tuple[dt.datetime, float, str]]:
    """(local time, E. coli per 100 ml, its "<" / ">" qualifier or "") for Alewife Brook inside the window."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb["Mystic Bacteria all yrs"].iter_rows(values_only=True))
    hdr = [str(c) for c in rows[5]]
    ei = hdr.index("E. coli (#/100mL)")
    qi = next(i for i, h in enumerate(hdr) if h.startswith("E. coli") and "'<'" in h)
    out = []
    for r in rows[6:]:
        if r[2] != "ALEWIFE BROOK" or r[ei] in ("", None):
            continue
        t = r[5] if isinstance(r[5], dt.datetime) else dt.datetime.fromisoformat(str(r[5]))
        if ALEWIFE_WINDOW[0] <= t <= ALEWIFE_WINDOW[1]:
            op = (r[qi] or "").strip() if isinstance(r[qi], str) or r[qi] is None else str(r[qi])
            if op not in ("", "<", ">"):
                raise RuntimeError(f"Alewife Brook {t}: unknown E. coli qualifier {op!r}")
            out.append((t, float(r[ei]), op))
    return out


def alewife(f: Fetcher) -> dict[str, Any]:
    """CSO starts (MassDEP portal + MWRA outfall MWR003) against MWRA E. coli; local (Eastern) times throughout."""
    md = f.fetch("massdep_alewife_incidents.xlsx", config.MASSDEP_ALEWIFE_XLSX, config.UNCLEAR,
                 headers=MASSDEP_HEADERS, restricted=True)
    mw = f.fetch("mwra_cso_table_export.csv", config.MWRA_CSO_CSV, config.UNCLEAR, headers=BROWSER, restricted=True)
    my = f.fetch("mwra_mystic_bacteria.xlsx", config.MWRA_MYSTIC_XLSX, config.UNCLEAR, headers=BROWSER,
                 restricted=True)
    starts = sorted(t for _, t in _massdep_starts(md) | _mwra_starts(mw.read_text(encoding="utf-8")))
    raw = _mystic_samples(my)
    smp = [(t, v, y) for t, v, op in raw if (y := config.label(v, op)) is not None]

    def hit(t: dt.datetime) -> bool:
        return any(0 <= (t - c).total_seconds() <= 48 * 3600 for c in starts)

    warned = [(v, y) for t, v, y in smp if hit(t)]
    quiet = [(v, y) for t, v, y in smp if not hit(t)]
    return {"cso_starts": len(starts), "samples": len(smp), "sampling_days": len({t.date() for t, _, _ in smp}),
            "warned": len(warned), "warned_exceed": sum(y for _, y in warned),
            "warned_median": st.median(v for v, _ in warned) if warned else None,
            "not_warned": len(quiet), "not_warned_exceed": sum(y for _, y in quiet),
            "not_warned_median": st.median(v for v, _ in quiet) if quiet else None,
            "warned_days": len({t.date() for t, _, _ in smp if hit(t)}),
            "qualified": sum(op != "" for _, _, op in raw), "ambiguous_dropped": len(raw) - len(smp)}
